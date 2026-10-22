"""Exportación de resultados: CSV, Excel, JSON, NumPy y FITS.

Las gráficas (PNG, SVG, PDF) se dibujan con Qt en ``turbulens.charts`` y se
exportan desde ``turbulens.reports``; este módulo no depende de la interfaz.
"""

from __future__ import annotations

import csv
import json
import math
import os
from typing import Iterable

import numpy as np

from .dimm import DimmParams, DimmResult, quality_label
from .fits import write_table

DATA_FORMATS = {
    "csv": "CSV",
    "xlsx": "Excel (.xlsx)",
    "json": "JSON",
    "npz": "NumPy (.npz)",
    "fits": "Tabla FITS",
}
CHART_FORMATS = {"pdfg": "PDF solo gráficas", "png": "PNG", "svg": "SVG"}
REPORT_FORMATS = {"pdf": "PDF completo"}
ALL_FORMATS = {**DATA_FORMATS, **CHART_FORMATS, **REPORT_FORMATS}

SUMMARY_LABELS = [
    ("video", "Video"),
    ("fecha_analisis", "Fecha del análisis"),
    ("tramo", "Tramo (cuadros)"),
    ("n_frames", "Cuadros analizados"),
    ("n_detected", "Cuadros con dos spots"),
    ("n_used", "Cuadros usados"),
    ("mean_sep_px", "Separación media (px)"),
    ("median_sep_px", "Mediana de la separación (px)"),
    ("std_sep_px", "Desviación estándar de la separación (px)"),
    ("sigma_l_px", "σ longitudinal (px)"),
    ("sigma_t_px", "σ transversal (px)"),
    ("eps_l", "Seeing longitudinal (″)"),
    ("eps_t", "Seeing transversal (″)"),
    ("eps", "Seeing (″)"),
    ("eps_zenith", "Seeing en el cenit (″)"),
    ("r0_cm", "Parámetro de Fried r0 (cm)"),
    ("calidad", "Calidad del cielo"),
    ("D_cm", "Diámetro de orificio (cm)"),
    ("B_cm", "Separación de orificios (cm)"),
    ("lambda_nm", "Longitud de onda (nm)"),
    ("scale_arcsec_px", "Escala de placa (″/px)"),
    ("altitude_deg", "Altura de la estrella (°)"),
]


def _clean(v):
    if isinstance(v, (np.floating, float)):
        v = float(v)
        return None if not math.isfinite(v) else v
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    return v


def summary_record(res, dimm: DimmResult, params: DimmParams) -> dict:
    import datetime as _dt

    d = {k: _clean(v) for k, v in dimm.summary_dict().items()}
    d.update({k: _clean(v) for k, v in params.to_dict().items()})
    d["calidad"] = quality_label(dimm.eps)
    d["video"] = res.video.get("name") if isinstance(res.video, dict) else ""
    d["tramo"] = f"{res.options.start}–{res.options.stop}"
    d["fecha_analisis"] = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    d["opciones"] = res.options.to_dict()
    d["info_video"] = {k: _clean(v) for k, v in (res.video or {}).items()}
    d["notas"] = list(res.notes)
    return d


def per_frame_columns(meas: np.ndarray, dimm: DimmResult | None) -> dict:
    cols = {
        "cuadro": meas["frame"],
        "tiempo_s": meas["t"],
        "x1_px": meas["x1"], "y1_px": meas["y1"],
        "x2_px": meas["x2"], "y2_px": meas["y2"],
        "flujo1": meas["flux1"].astype(np.float64), "flujo2": meas["flux2"].astype(np.float64),
        "region_x": meas["roi_x"], "region_y": meas["roi_y"],
        "detectado": meas["valid"],
    }
    if dimm is not None and dimm.arrays:
        a = dimm.arrays
        cols["separacion_px"] = a["sep"]
        cols["longitudinal_px"] = a["lon"]
        cols["transversal_px"] = a["tra"]
        cols["usado"] = a["used"]
    return cols


def write_measurements_csv(path: str, meas: np.ndarray, dimm: DimmResult | None) -> None:
    cols = per_frame_columns(meas, dimm)
    names = list(cols)
    arrs = [np.asarray(cols[n]) for n in names]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(names)
        for i in range(len(meas)):
            row = []
            for a in arrs:
                v = a[i]
                if a.dtype == bool:
                    row.append(int(v))
                elif np.issubdtype(a.dtype, np.floating):
                    row.append("" if not np.isfinite(v) else f"{v:.5f}")
                else:
                    row.append(int(v))
            w.writerow(row)


def _summary_rows(rec: dict) -> list[tuple[str, object]]:
    rows = []
    for key, label in SUMMARY_LABELS:
        if key in rec:
            v = rec[key]
            if isinstance(v, float):
                v = round(v, 4)
            rows.append((label, "" if v is None else v))
    return rows


def export_analysis(rec: dict, meas: np.ndarray, dimm: DimmResult | None, formats: Iterable[str], folder: str,
                    base: str = "turbulens") -> list[str]:
    """Exporta un análisis en los formatos de datos pedidos. Devuelve las rutas creadas."""
    os.makedirs(folder, exist_ok=True)
    out = []
    formats = set(formats)
    cols = per_frame_columns(meas, dimm)
    if "csv" in formats:
        p = os.path.join(folder, f"{base}_mediciones.csv")
        write_measurements_csv(p, meas, dimm)
        p2 = os.path.join(folder, f"{base}_resumen.csv")
        with open(p2, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["campo", "valor"])
            w.writerows(_summary_rows(rec))
        out += [p, p2]
    if "json" in formats:
        p = os.path.join(folder, f"{base}.json")
        data = {"resumen": rec, "cuadros": {k: [_clean(v) for v in np.asarray(a).tolist()] for k, a in cols.items()}}
        with open(p, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1, default=_clean)
        out.append(p)
    if "npz" in formats:
        p = os.path.join(folder, f"{base}.npz")
        np.savez_compressed(p, mediciones=meas, **{k: np.asarray(v) for k, v in cols.items()},
                            resumen=json.dumps(rec, ensure_ascii=False, default=_clean))
        out.append(p)
    if "fits" in formats:
        p = os.path.join(folder, f"{base}_mediciones.fits")
        extra = {"SEEING": rec.get("eps"), "SIGL_PX": rec.get("sigma_l_px"), "SIGT_PX": rec.get("sigma_t_px"),
                 "SCALE": rec.get("scale_arcsec_px"), "D_CM": rec.get("D_cm"), "B_CM": rec.get("B_cm"),
                 "LAMBDA": rec.get("lambda_nm")}
        write_table(p, cols, units={"tiempo_s": "s", "x1_px": "pix", "y1_px": "pix", "x2_px": "pix", "y2_px": "pix",
                                    "separacion_px": "pix", "longitudinal_px": "pix", "transversal_px": "pix"},
                    extra={k: v for k, v in extra.items() if v is not None})
        out.append(p)
    if "xlsx" in formats:
        p = os.path.join(folder, f"{base}.xlsx")
        write_xlsx(p, [("Resumen", ["Campo", "Valor"], _summary_rows(rec)),
                       ("Cuadros", list(cols), _rows_from_cols(cols))])
        out.append(p)
    return out


def _rows_from_cols(cols: dict) -> list[list]:
    arrs = [np.asarray(a) for a in cols.values()]
    rows = []
    for i in range(len(arrs[0]) if arrs else 0):
        r = []
        for a in arrs:
            v = a[i]
            if a.dtype == bool:
                r.append(int(v))
            elif np.issubdtype(a.dtype, np.floating):
                r.append(None if not np.isfinite(v) else float(v))
            else:
                r.append(int(v))
        rows.append(r)
    return rows


def write_xlsx(path: str, sheets: list) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(header, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(str(h)) + 4))
    wb.save(path)


HISTORY_COLUMNS = [
    ("id", "ID"), ("created", "Guardado"), ("label", "Etiqueta"), ("video_name", "Video"),
    ("start", "Cuadro inicial"), ("stop", "Cuadro final"), ("n_used", "Cuadros usados"),
    ("mean_sep_px", "Separación media (px)"), ("std_sep_px", "σ separación (px)"),
    ("sigma_l_px", "σ longitudinal (px)"), ("sigma_t_px", "σ transversal (px)"),
    ("eps", "Seeing (″)"), ("eps_zenith", "Seeing cenit (″)"), ("r0_cm", "r0 (cm)"),
    ("D_cm", "D (cm)"), ("B_cm", "B (cm)"), ("lambda_nm", "λ (nm)"), ("scale", "Escala (″/px)"),
    ("altitude", "Altura (°)"), ("notes", "Notas"),
]


def export_history_table(records: list[dict], formats: Iterable[str], folder: str, base: str = "historial") -> list[str]:
    os.makedirs(folder, exist_ok=True)
    formats = set(formats)
    header = [lab for _, lab in HISTORY_COLUMNS]
    rows = [[_clean(r.get(k)) for k, _ in HISTORY_COLUMNS] for r in records]
    out = []
    if "csv" in formats:
        p = os.path.join(folder, f"{base}.csv")
        with open(p, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows([["" if v is None else v for v in r] for r in rows])
        out.append(p)
    if "xlsx" in formats:
        p = os.path.join(folder, f"{base}.xlsx")
        write_xlsx(p, [("Historial", header, rows)])
        out.append(p)
    if "json" in formats:
        p = os.path.join(folder, f"{base}.json")
        with open(p, "w", encoding="utf-8") as f:
            json.dump([{k: _clean(r.get(k)) for k, _ in HISTORY_COLUMNS} for r in records], f, ensure_ascii=False, indent=1)
        out.append(p)
    if "npz" in formats:
        p = os.path.join(folder, f"{base}.npz")
        num = {k: np.array([np.nan if r.get(k) is None else r.get(k) for r in records], dtype=float)
               for k, _ in HISTORY_COLUMNS if k not in ("created", "label", "video_name", "notes")}
        np.savez_compressed(p, **num)
        out.append(p)
    if "fits" in formats:
        p = os.path.join(folder, f"{base}.fits")
        num = {k: np.array([np.nan if r.get(k) is None else r.get(k) for r in records], dtype=float)
               for k, _ in HISTORY_COLUMNS if k not in ("created", "label", "video_name", "notes")}
        write_table(p, num)
        out.append(p)
    return out
