"""Pruebas del motor de análisis (no requieren interfaz gráfica)."""

import json
import math
import os

import numpy as np
import pytest

from turbulens.core import dimm, export, fits, pipeline, stats
from turbulens.core.detect import config_for_radius, detect_pairs, estimate_spot_size
from turbulens.core.history import History

from synthetic import dimm_track, make_video, render_pair


def test_detect_subpixel_accuracy():
    rng = np.random.default_rng(5)
    n = 200
    p1 = np.c_[30 + rng.uniform(-2, 2, n), 28 + rng.uniform(-2, 2, n)]
    p2 = p1 + np.c_[rng.uniform(5, 8, n), rng.uniform(9, 11, n)]
    stack = np.stack([render_pair(64, 64, p1[i], p2[i], rng=rng) for i in range(n)])
    d = detect_pairs(stack, config_for_radius(estimate_spot_size(stack)))
    assert d["valid"].all()
    got1 = np.c_[d["x1"], d["y1"]]
    got2 = np.c_[d["x2"], d["y2"]]
    # El orden de los spots puede variar: emparejar por cercanía.
    swap = np.linalg.norm(got1 - p1, axis=1) > np.linalg.norm(got1 - p2, axis=1)
    g1 = np.where(swap[:, None], got2, got1)
    g2 = np.where(swap[:, None], got1, got2)
    err = np.concatenate([np.linalg.norm(g1 - p1, axis=1), np.linalg.norm(g2 - p2, axis=1)])
    assert np.median(err) < 0.06
    assert err.max() < 0.3


def test_single_spot_is_rejected():
    img = render_pair(64, 64, (32, 32), (32.0, 32.0), peak=200)
    d = detect_pairs(img[None])
    assert not d["valid"][0]


def test_tokovinin_coefficients():
    kl, kt = dimm.k_coefficients(14 / 2.5)
    assert kl == pytest.approx(0.2549, abs=2e-3)
    assert kt == pytest.approx(0.2006, abs=2e-3)


def test_seeing_reference_value():
    # σ = 0,72 px con 1″/px, D = 2,5 cm, B = 14 cm, λ = 550 nm  →  ε ≈ 1,11″ (cálculo de referencia del proyecto)
    kl, _ = dimm.k_coefficients(14 / 2.5)
    r0 = dimm.r0_from_variance((0.72 / dimm.ARCSEC_PER_RAD) ** 2, kl, 550e-9, 0.025)
    eps = dimm.seeing_from_r0(r0, 550e-9)
    assert eps == pytest.approx(1.11, abs=0.02)
    assert r0 * 100 == pytest.approx(10.0, abs=0.3)


def test_analyze_recovers_sigma_and_cancels_common_motion():
    n = 6000
    p1, p2 = dimm_track(n, sigma_l=0.7, sigma_t=0.5)
    meas = np.zeros(n, dtype=pipeline.MEAS_DTYPE)
    meas["x1"], meas["y1"] = p1[:, 0], p1[:, 1]
    meas["x2"], meas["y2"] = p2[:, 0], p2[:, 1]
    meas["valid"] = True
    # Intercambiar spots en la mitad de los cuadros no debe afectar
    sw = np.arange(n) % 2 == 0
    for a, b in (("x1", "x2"), ("y1", "y2")):
        t = meas[a][sw].copy()
        meas[a][sw] = meas[b][sw]
        meas[b][sw] = t
    r = dimm.analyze(meas, dimm.DimmParams(scale_arcsec_px=1.0))
    assert r.sigma_l_px == pytest.approx(0.7, rel=0.05)
    assert r.sigma_t_px == pytest.approx(0.5, rel=0.05)
    assert r.mean_sep_px == pytest.approx(13.0, abs=0.05)
    assert math.isfinite(r.eps) and r.eps > 0


def test_outliers_are_clipped():
    rng = np.random.default_rng(2)
    x = np.concatenate([rng.normal(13, 0.7, 5000), rng.uniform(0, 1, 60)])
    keep = stats.sigma_clip_mask(x, 5)
    assert (~keep).sum() >= 60
    assert np.std(x[keep]) == pytest.approx(0.7, rel=0.05)


def test_running_std_and_summary():
    t = np.arange(600) / 60
    x = np.r_[np.ones(300), np.ones(300) * 2] + np.tile([0.1, -0.1], 300)
    mids, sd, cnt = stats.running_std(t, x, 1.0)
    assert len(mids) == 10 and np.allclose(cnt, 60)
    s = stats.summarize([1, 2, 3, 4, 100])
    assert s.median == 3 and s.n == 5


def test_fits_cube_and_table(tmp_path):
    p = tmp_path / "cubo.fits"
    w = fits.FitsCubeWriter(str(p), 5, 8, 6, np.uint16)
    for i in range(3):
        w.write(np.full((8, 6), 1000 + i, np.uint16))
    w.close()
    raw = p.read_bytes()
    assert len(raw) % 2880 == 0
    hdr = raw[:2880].decode("ascii")
    assert "NAXIS3  =                    3" in hdr
    assert hdr.startswith("SIMPLE  =                    T")
    data = np.frombuffer(raw[2880:2880 + 8 * 6 * 2], dtype=">i2").astype(int) + 32768
    assert (data == 1000).all()
    t = tmp_path / "tabla.fits"
    fits.write_table(str(t), {"a": np.arange(4.0), "b": np.array([1, 2, 3, 4]), "c": np.array([True, False, True, True])})
    raw = t.read_bytes()
    assert len(raw) % 2880 == 0
    assert "BINTABLE" in raw[2880:5760].decode("ascii")


@pytest.mark.parametrize("fmt", ["PNG", "TIFF", "JPG", "NPY", "FITS"])
def test_pipeline_on_synthetic_video(tmp_path, fmt):
    av = pytest.importorskip("av")
    vid = tmp_path / "sintetico.mkv"
    n = 240
    p1, p2 = make_video(str(vid), n=n, sigma_l=0.7, sigma_t=0.6)
    from turbulens.core.video import VideoReader

    r = VideoReader(str(vid))
    assert r.info.n_frames == n
    roi = pipeline.fit_roi(r.frame_at(0))
    opts = pipeline.ProcessOptions(roi=roi, start=0, stop=n - 1, out_dir=str(tmp_path / "out"), fmt=fmt, every_n=2)
    res = pipeline.run(r, opts)
    assert res.meas.size == n
    assert res.meas["valid"].all(), "la estrella debe seguirse en todos los cuadros pese a la deriva"
    assert res.n_saved == n // 2
    d = dimm.analyze(res.meas, dimm.DimmParams(scale_arcsec_px=1.0))
    assert d.sigma_l_px == pytest.approx(0.7, rel=0.15)
    assert d.mean_sep_px == pytest.approx(13.0, abs=0.1)
    files = os.listdir(tmp_path / "out")
    if fmt in ("NPY", "FITS"):
        assert len(files) == 1
        if fmt == "NPY":
            cube = np.load(tmp_path / "out" / files[0])
            assert cube.shape == (n // 2, roi[3], roi[2])
    else:
        assert len(files) == n // 2
    pipeline.write_outputs(res, d, dimm.DimmParams(scale_arcsec_px=1.0))
    assert (tmp_path / "out" / "mediciones.csv").exists()
    summ = json.loads((tmp_path / "out" / "resumen.json").read_text(encoding="utf-8"))
    assert summ["n_frames"] == n


def test_exports_and_history(tmp_path):
    n = 500
    p1, p2 = dimm_track(n)
    meas = np.zeros(n, dtype=pipeline.MEAS_DTYPE)
    meas["frame"] = np.arange(n)
    meas["t"] = np.arange(n) / 60
    meas["x1"], meas["y1"] = p1[:, 0], p1[:, 1]
    meas["x2"], meas["y2"] = p2[:, 0], p2[:, 1]
    meas["valid"] = True
    prm = dimm.DimmParams(scale_arcsec_px=1.0, altitude_deg=60)
    d = dimm.analyze(meas, prm)
    opts = pipeline.ProcessOptions(roi=(0, 0, 64, 64), start=0, stop=n - 1)
    res = pipeline.ProcessResult(meas=meas, options=opts, video={"name": "x.mov", "fps": 60}, n_saved=0, elapsed_s=1,
                                 cancelled=False, out_dir="", detect_cfg={})
    rec = export.summary_record(res, d, prm)
    out = export.export_analysis(rec, meas, d, ["csv", "xlsx", "json", "npz", "fits"], str(tmp_path), "prueba")
    assert len(out) == 6 and all(os.path.getsize(p) > 0 for p in out)
    from openpyxl import load_workbook

    wb = load_workbook(tmp_path / "prueba.xlsx")
    assert set(wb.sheetnames) == {"Resumen", "Cuadros"}
    assert wb["Cuadros"].max_row == n + 1

    h = History("analisis", root=str(tmp_path / "datos"))
    he = History("educativo", root=str(tmp_path / "datos"))
    rid = h.add({"label": "noche 1", "eps": d.eps, "std_sep_px": d.std_sep_px, "roi": [0, 0, 64, 64]}, meas=meas, arrays=d.arrays)
    assert he.list() == []  # modos separados
    assert h.list()[0]["label"] == "noche 1"
    assert h.load_arrays(rid)["mediciones"].size == n
    files = export.export_history_table(h.list(), ["csv", "xlsx", "json", "npz", "fits"], str(tmp_path / "hist"))
    assert len(files) == 5
    h.delete(rid)
    assert h.list() == []
